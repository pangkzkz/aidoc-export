# -*- coding: utf-8 -*-
"""aidoc_export - Convert LLM answers (Markdown) into DOCX / PDF reports.

Example::

    from aidoc_export import ExportConfig, export_docx, export_pdf

    cfg = ExportConfig(title="Weekly Report", author="Data Team", font_path="fonts/NanumGothic.ttf")
    export_docx(markdown_text, "weekly.docx", cfg)
    export_pdf(markdown_text, "weekly.pdf", cfg)

Structure:
    1. parse_markdown()  : markdown-it-py tokens -> list of internal Blocks
    2. _DocxWriter       : Blocks -> python-docx
    3. _PdfWriter        : Blocks -> reportlab (platypus)
"""
from __future__ import annotations

import argparse
import io
import logging
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, List, Optional, Sequence, Tuple, Union
from xml.sax.saxutils import escape

from markdown_it import MarkdownIt

__version__ = "0.3.0"

log = logging.getLogger("aidoc_export")

PathOrFile = Union[str, Path, IO[bytes], None]


# ===========================================================================
# Configuration
# ===========================================================================

@dataclass
class Source:
    """An entry in the 'Sources' list at the end of the report."""

    title: str
    url: Optional[str] = None
    page: Optional[Union[int, str]] = None


@dataclass
class ExportConfig:
    """Export options.

    Attributes:
        title: Document title (file properties, used in the header)
        author: Author (file properties)
        page_size: "A4" or "Letter"
        margin_mm: Page margin (mm)
        base_size: Body font size (pt)
        header: Whether to show a header. If True, prints the title
        page_numbers: Whether to show page numbers ("3 / 12") in the footer
        sources: List of sources. Nothing is printed if empty
        font_path: TTF path for PDF body text. If unset, searched via env vars / system
        bold_font_path: TTF path for PDF bold font (optional)
        docx_font: Font name written to the DOCX (not embedded in the file)
        sources_title: Heading for the sources list
    """

    title: Optional[str] = None
    author: Optional[str] = None
    page_size: str = "A4"
    margin_mm: float = 20.0
    base_size: float = 10.5
    header: bool = False
    page_numbers: bool = True
    sources: List[Source] = field(default_factory=list)
    font_path: Optional[str] = None
    bold_font_path: Optional[str] = None
    docx_font: str = "Malgun Gothic"
    sources_title: str = "Sources"


class FontNotFoundError(RuntimeError):
    """No Korean TTF font could be found for the PDF."""


# ===========================================================================
# 1. Markdown -> Blocks
# ===========================================================================

@dataclass
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    underline: bool = False
    code: bool = False
    mark: bool = False
    link: Optional[str] = None
    br: bool = False


@dataclass
class Block:
    """kind: heading | para | list | table | code | quote | hr | pagebreak

    - heading: level, spans
    - para: spans
    - list: ordered, start, items (each item = (checked, List[Block]))
    - table: rows (List[List[List[Span]]]), aligns, header_rows
    - code: text
    - quote: children
    """

    kind: str
    spans: List[Span] = field(default_factory=list)
    level: int = 0
    ordered: bool = False
    start: int = 1
    items: List[Tuple[Optional[bool], List["Block"]]] = field(default_factory=list)
    rows: List[List[List[Span]]] = field(default_factory=list)
    aligns: List[Optional[str]] = field(default_factory=list)
    header_rows: int = 0
    text: str = ""
    children: List["Block"] = field(default_factory=list)


_PAGEBREAK = re.compile(r"^\s*(<!--\s*page-?break\s*-->|\\pagebreak|\\newpage)\s*$", re.I)
_TASK = re.compile(r"^\[([ xX])\]\s+")
_MARK = re.compile(r"==([^=\n]+)==")

_md = MarkdownIt("commonmark", {"html": True}).enable(["table", "strikethrough"])


def parse_markdown(text: str) -> List[Block]:
    """Convert a Markdown string into a list of Blocks."""
    tokens = _md.parse(text.replace("\r\n", "\n"))
    blocks, _ = _parse_blocks(tokens, 0, None)
    return blocks


def _parse_blocks(tokens, i, stop_type):
    out: List[Block] = []
    while i < len(tokens):
        t = tokens[i]
        if stop_type and t.type == stop_type:
            return out, i + 1

        if t.type == "heading_open":
            out.append(Block("heading", level=int(t.tag[1]), spans=_inline(tokens[i + 1])))
            i += 3
        elif t.type == "paragraph_open":
            inline = tokens[i + 1]
            if _PAGEBREAK.match(inline.content):
                out.append(Block("pagebreak"))
            else:
                out.append(Block("para", spans=_inline(inline)))
            i += 3
        elif t.type in ("bullet_list_open", "ordered_list_open"):
            ordered = t.type == "ordered_list_open"
            start = int(t.attrGet("start") or 1)
            close = "ordered_list_close" if ordered else "bullet_list_close"
            items = []
            i += 1
            while tokens[i].type != close:
                # list_item_open
                children, i = _parse_blocks(tokens, i + 1, "list_item_close")
                checked = None
                if children and children[0].kind == "para" and children[0].spans:
                    first = children[0].spans[0]
                    m = _TASK.match(first.text)
                    if m:
                        checked = m.group(1).lower() == "x"
                        first.text = first.text[m.end():]
                items.append((checked, children))
            out.append(Block("list", ordered=ordered, start=start, items=items))
            i += 1
        elif t.type == "table_open":
            block, i = _parse_table(tokens, i + 1)
            out.append(block)
        elif t.type in ("fence", "code_block"):
            out.append(Block("code", text=t.content.rstrip("\n")))
            i += 1
        elif t.type == "blockquote_open":
            children, i = _parse_blocks(tokens, i + 1, "blockquote_close")
            out.append(Block("quote", children=children))
        elif t.type == "hr":
            out.append(Block("hr"))
            i += 1
        elif t.type == "html_block":
            if _PAGEBREAK.match(t.content.strip()):
                out.append(Block("pagebreak"))
            else:
                log.debug("Ignoring HTML block: %r", t.content[:40])
            i += 1
        else:
            i += 1
    return out, i


def _parse_table(tokens, i):
    rows: List[List[List[Span]]] = []
    aligns: List[Optional[str]] = []
    header_rows = 0
    in_head = False
    while tokens[i].type != "table_close":
        t = tokens[i]
        if t.type == "thead_open":
            in_head = True
        elif t.type == "thead_close":
            in_head = False
        elif t.type == "tr_open":
            rows.append([])
            if in_head:
                header_rows += 1
        elif t.type in ("th_open", "td_open"):
            style = t.attrGet("style") or ""
            m = re.search(r"text-align:(\w+)", style)
            if len(rows) == 1:
                aligns.append(m.group(1) if m else None)
            rows[-1].append(_inline(tokens[i + 1]))
            i += 2
        i += 1
    return Block("table", rows=rows, aligns=aligns, header_rows=header_rows), i + 1


def _inline(token) -> List[Span]:
    spans: List[Span] = []
    st = {"bold": False, "italic": False, "strike": False, "underline": False, "mark": False, "link": None}

    def add(text, **extra):
        spans.append(Span(text=text, **{**st, **extra}))

    for c in token.children or []:
        tp = c.type
        if tp == "text":
            pos = 0
            for m in _MARK.finditer(c.content):  # ==highlight==
                if m.start() > pos:
                    add(c.content[pos:m.start()])
                add(m.group(1), mark=True)
                pos = m.end()
            if pos < len(c.content):
                add(c.content[pos:])
        elif tp == "softbreak":
            add(" ")
        elif tp == "hardbreak":
            add("", br=True)
        elif tp == "code_inline":
            add(c.content, code=True)
        elif tp in ("strong_open", "strong_close"):
            st["bold"] = tp == "strong_open"
        elif tp in ("em_open", "em_close"):
            st["italic"] = tp == "em_open"
        elif tp in ("s_open", "s_close"):
            st["strike"] = tp == "s_open"
        elif tp == "link_open":
            st["link"] = c.attrGet("href")
        elif tp == "link_close":
            st["link"] = None
        elif tp == "image":
            add("[%s]" % (c.content or "image"))
        elif tp == "html_inline":
            tag = c.content.strip().lower()
            if re.match(r"<br\s*/?>", tag):
                add("", br=True)
            elif tag in ("<mark>", "</mark>"):
                st["mark"] = tag == "<mark>"
            elif tag in ("<u>", "</u>"):
                st["underline"] = tag == "<u>"
    return [s for s in spans if s.text or s.br]


def _append_sources(blocks: List[Block], cfg: ExportConfig) -> List[Block]:
    if not cfg.sources:
        return blocks
    items = []
    for s in cfg.sources:
        spans = [Span(s.title)]
        if s.page not in (None, ""):
            spans.append(Span(", %s%s" % ("p. " if isinstance(s.page, int) else "", s.page)))
        if s.url:
            spans += [Span(" - "), Span(s.url, link=s.url)]
        items.append((None, [Block("para", spans=spans)]))
    return blocks + [Block("heading", level=2, spans=[Span(cfg.sources_title)]),
                     Block("list", ordered=True, items=items)]


def _plain(spans: Sequence[Span]) -> str:
    return "".join("\n" if s.br else s.text for s in spans)


# ===========================================================================
# 2. DOCX
# ===========================================================================

class _DocxWriter:
    MONO = "Consolas"

    def __init__(self, cfg: ExportConfig):
        from docx import Document
        from docx.shared import Mm, Pt

        self.cfg = cfg
        self.doc = Document()
        self.Pt = Pt
        sec = self.doc.sections[0]
        if cfg.page_size.upper() == "LETTER":
            sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
        else:
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Mm(cfg.margin_mm))

        normal = self.doc.styles["Normal"]
        normal.font.size = Pt(cfg.base_size)
        self._set_font(normal.element, cfg.docx_font)
        normal.paragraph_format.space_after = Pt(cfg.base_size * 0.6)
        normal.paragraph_format.line_spacing = 1.3
        for lv, scale in zip(range(1, 7), (1.75, 1.45, 1.22, 1.1, 1.0, 1.0)):
            h = self.doc.styles["Heading %d" % lv]
            h.font.size = Pt(cfg.base_size * scale)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = None
            self._set_font(h.element, cfg.docx_font)

        props = self.doc.core_properties
        props.title = cfg.title or ""
        props.author = cfg.author or ""
        props.language = "ko-KR"
        self._header_footer()

    # ---- Low-level XML helpers ---------------------------------------------

    @staticmethod
    def _set_font(el, name: str, latin: Optional[str] = None) -> None:
        """Also set the East Asian (eastAsia) font. python-docx's font.name only changes ascii/hAnsi."""
        from docx.oxml.ns import qn

        rpr = el.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), latin or name)
        fonts.set(qn("w:eastAsia"), name)

    @staticmethod
    def _shade(el, fill: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        el.append(shd)

    @staticmethod
    def _field(paragraph, instr: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), instr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        paragraph._p.append(fld)

    def _header_footer(self) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import RGBColor

        sec = self.doc.sections[0]
        if self.cfg.header and self.cfg.title:
            p = sec.header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(self.cfg.title)
            run.font.size = self.Pt(8)
            run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        if self.cfg.page_numbers:
            p = sec.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._field(p, "PAGE")
            p.add_run(" / ")
            self._field(p, "NUMPAGES")

    # ---- Blocks -------------------------------------------------------------

    def write(self, blocks: List[Block]) -> None:
        for b in blocks:
            self._block(b, indent=0, quote=False)

    def _block(self, b: Block, indent: int, quote: bool, container=None) -> None:
        from docx.shared import Mm, RGBColor

        target = container or self.doc
        if b.kind == "heading":
            p = target.add_paragraph(style="Heading %d" % min(b.level, 6))
            p.paragraph_format.keep_with_next = True
            self._runs(p, b.spans)
        elif b.kind == "para":
            p = target.add_paragraph()
            self._runs(p, b.spans)
            if quote:
                p.paragraph_format.left_indent = Mm(6)
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0x5F, 0x5B, 0x55)
        elif b.kind == "list":
            for n, (checked, children) in enumerate(b.items, start=b.start):
                if b.ordered:
                    marker = "%d. " % n
                elif checked is not None:
                    marker = "■ " if checked else "□ "
                else:
                    marker = ("•", "◦", "▪")[min(indent, 2)] + " "
                for k, child in enumerate(children):
                    if child.kind == "list":
                        self._block(child, indent + 1, quote, container)
                        continue
                    if child.kind != "para":
                        self._block(child, indent, quote, container)
                        continue
                    p = target.add_paragraph()
                    pf = p.paragraph_format
                    pf.left_indent = Mm(7 * (indent + 1))
                    pf.first_line_indent = Mm(-5) if k == 0 else None
                    pf.space_after = self.Pt(2)
                    if k == 0:
                        p.add_run(marker)
                    self._runs(p, child.spans)
        elif b.kind == "table":
            self._table(b, target)
        elif b.kind == "code":
            for line in b.text.split("\n") or [""]:
                p = target.add_paragraph()
                p.paragraph_format.space_after = self.Pt(0)
                self._shade(p._p.get_or_add_pPr(), "F3F2EF")
                r = p.add_run(line or " ")
                self._set_font(r._r, self.cfg.docx_font, latin=self.MONO)
                r.font.size = self.Pt(self.cfg.base_size - 1)
            target.add_paragraph().paragraph_format.space_after = self.Pt(0)
        elif b.kind == "quote":
            for child in b.children:
                self._block(child, indent, True, container)
        elif b.kind == "hr":
            self._hr(target.add_paragraph())
        elif b.kind == "pagebreak":
            from docx.enum.text import WD_BREAK

            target.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _runs(self, p, spans: Sequence[Span]) -> None:
        from docx.enum.text import WD_COLOR_INDEX
        from docx.shared import RGBColor

        for s in spans:
            if s.br:
                p.add_run().add_break()
                continue
            if s.link:
                self._hyperlink(p, s)
                continue
            r = p.add_run(s.text)
            r.bold = s.bold or None
            r.italic = s.italic or None
            r.font.strike = s.strike or None
            r.underline = s.underline or None
            if s.mark:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if s.code:
                self._set_font(r._r, self.cfg.docx_font, latin=self.MONO)
                r.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)

    def _hyperlink(self, p, s: Span) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        rid = p.part.relate_to(s.link, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        for tag, val in (("w:color", "1F5BD8"), ("w:u", "single")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            rpr.append(e)
        if s.bold:
            rpr.append(OxmlElement("w:b"))
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = s.text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def _hr(self, p) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "CCCCCC")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.append(bdr)

    def _table(self, b: Block, target) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        ncols = max(len(r) for r in b.rows)
        table = target.add_table(rows=len(b.rows), cols=ncols)
        table.style = "Table Grid"
        align_map = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for ri, row in enumerate(b.rows):
            is_head = ri < b.header_rows
            if is_head:  # repeat header rows across page breaks
                trpr = table.rows[ri]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:tblHeader"))
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                spans = row[ci] if ci < len(row) else []
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = self.Pt(0)
                align = b.aligns[ci] if ci < len(b.aligns) else None
                if align in align_map:
                    p.alignment = align_map[align]
                if is_head:
                    spans = [Span(**{**s.__dict__, "bold": True}) for s in spans]
                    self._shade(cell._tc.get_or_add_tcPr(), "EFEEEA")
                self._runs(p, spans)
        target.add_paragraph().paragraph_format.space_after = self.Pt(0)

    def save(self, out: PathOrFile) -> Optional[bytes]:
        if out is None:
            buf = io.BytesIO()
            self.doc.save(buf)
            return buf.getvalue()
        self.doc.save(out if not isinstance(out, Path) else str(out))
        return None


# ===========================================================================
# 3. PDF
# ===========================================================================

_FONT_CANDIDATES = [
    # (regular, bold)
    ("C:/Windows/Fonts/malgun.ttf", "C:/Windows/Fonts/malgunbd.ttf"),
    ("/usr/share/fonts/truetype/nanum/NanumGothic.ttf", "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
    ("/usr/share/fonts/nanum/NanumGothic.ttf", "/usr/share/fonts/nanum/NanumGothicBold.ttf"),
    ("~/Library/Fonts/NanumGothic.ttf", "~/Library/Fonts/NanumGothicBold.ttf"),
    ("/Library/Fonts/NanumGothic.ttf", "/Library/Fonts/NanumGothicBold.ttf"),
    ("~/Library/Fonts/Pretendard-Regular.ttf", "~/Library/Fonts/Pretendard-Bold.ttf"),
    ("fonts/Pretendard-Regular.ttf", "fonts/Pretendard-Bold.ttf"),
    ("fonts/NanumGothic.ttf", "fonts/NanumGothicBold.ttf"),
]


def find_korean_font(cfg: ExportConfig) -> Tuple[str, Optional[str]]:
    """Find a Korean TTF for PDF. Order: config > env vars (AIDOC_FONT, AIDOC_FONT_BOLD) > known paths."""
    if cfg.font_path:
        return cfg.font_path, cfg.bold_font_path
    env = os.environ.get("AIDOC_FONT")
    if env:
        return env, os.environ.get("AIDOC_FONT_BOLD")
    for regular, bold in _FONT_CANDIDATES:
        r, b = os.path.expanduser(regular), os.path.expanduser(bold)
        if os.path.isfile(r):
            return r, (b if os.path.isfile(b) else None)
    raise FontNotFoundError(
        "Could not find a Korean TTF font for PDF. Set it via ExportConfig(font_path=...) or "
        "the AIDOC_FONT environment variable. (e.g. NanumGothic.ttf, Pretendard-Regular.ttf, malgun.ttf)"
    )


_registered = {}


def _register_fonts(cfg: ExportConfig) -> Tuple[str, str]:
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFError, TTFont

    regular, bold = find_korean_font(cfg)
    key = (regular, bold)
    if key in _registered:
        return _registered[key]
    n = len(_registered) + 1
    names = ("AidocKR%d" % n, "AidocKR%d-Bold" % n)
    try:
        pdfmetrics.registerFont(TTFont(names[0], regular))
        pdfmetrics.registerFont(TTFont(names[1], bold or regular))
    except TTFError as e:
        raise FontNotFoundError(
            "Cannot use %s: %s\n"
            "reportlab only supports TrueType (.ttf). CFF-based .otf / .ttc (e.g. Noto Sans CJK) will not work." % (regular, e)
        ) from e
    for (b_, i_), name in (((0, 0), names[0]), ((1, 0), names[1]), ((0, 1), names[0]), ((1, 1), names[1])):
        addMapping(names[0], b_, i_, name)
    _registered[key] = names
    log.info("PDF font: %s / %s", regular, bold or "(no bold)")
    return names


class _PdfWriter:
    def __init__(self, cfg: ExportConfig):
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle

        self.cfg = cfg
        self.font, self.bold = _register_fonts(cfg)
        size = cfg.base_size
        self.colors = colors
        base = ParagraphStyle("body", fontName=self.font, fontSize=size, leading=size * 1.6,
                              spaceAfter=size * 0.6, textColor=colors.HexColor("#1F1E1C"))
        self.st = {"body": base}
        for lv, scale in zip(range(1, 7), (1.8, 1.45, 1.22, 1.1, 1.0, 1.0)):
            self.st["h%d" % lv] = ParagraphStyle(
                "h%d" % lv, parent=base, fontName=self.bold, fontSize=size * scale, leading=size * scale * 1.35,
                spaceBefore=size * (1.1 if lv <= 2 else 0.8), spaceAfter=size * 0.5, keepWithNext=1)
        self.st["cell"] = ParagraphStyle("cell", parent=base, spaceAfter=0, leading=size * 1.45)
        self.st["quote"] = ParagraphStyle("quote", parent=base, textColor=colors.HexColor("#5F5B55"), spaceAfter=2)
        self.st["code"] = ParagraphStyle("code", parent=base, fontSize=size * 0.9, leading=size * 1.35)
        self.st["li"] = ParagraphStyle("li", parent=base, spaceAfter=size * 0.2)

    # ---- Inline -> reportlab mini-HTML --------------------------------------

    def _markup(self, spans: Sequence[Span]) -> str:
        out = []
        for s in spans:
            if s.br:
                out.append("<br/>")
                continue
            t = escape(s.text)
            if s.code:
                t = '<font backColor="#F1F0EC" color="#B03A2E">%s</font>' % t
            if s.mark:
                t = '<font backColor="#FFF1A8">%s</font>' % t
            if s.bold:
                t = "<b>%s</b>" % t
            if s.italic:
                t = "<i>%s</i>" % t
            if s.strike:
                t = "<strike>%s</strike>" % t
            if s.underline:
                t = "<u>%s</u>" % t
            if s.link:
                t = '<a href="%s" color="#1F5BD8"><u>%s</u></a>' % (escape(s.link, {'"': "&quot;"}), t)
            out.append(t)
        return "".join(out)

    def _cell_style(self, align: Optional[str]):
        if align not in ("center", "right"):
            return self.st["cell"]
        key = "cell_" + align
        if key not in self.st:
            self.st[key] = self.st["cell"].clone(key, alignment=1 if align == "center" else 2)
        return self.st[key]

    # ---- Blocks -------------------------------------------------------------

    def flowables(self, blocks: List[Block], width: float, style: str = "body") -> list:
        from reportlab.platypus import (HRFlowable, ListFlowable, ListItem, PageBreak, Paragraph,
                                        Preformatted, Spacer, Table, TableStyle)

        c = self.colors
        out = []
        for b in blocks:
            if b.kind == "heading":
                out.append(Paragraph(self._markup(b.spans), self.st["h%d" % min(b.level, 6)]))
            elif b.kind == "para":
                out.append(Paragraph(self._markup(b.spans), self.st[style]))
            elif b.kind == "list":
                items = []
                for checked, children in b.items:
                    if checked is not None and children and children[0].kind == "para":
                        children = [Block("para", spans=[Span("■ " if checked else "□ ")] + children[0].spans)] + children[1:]
                    inner = self.flowables(children, width - 18, "li")
                    items.append(ListItem(inner or [Paragraph("", self.st["li"])]))
                task = all(ch is not None for ch, _ in b.items)
                kw = dict(bulletFontName=self.font, bulletFontSize=self.cfg.base_size, leftIndent=16, spaceAfter=4)
                if b.ordered:
                    out.append(ListFlowable(items, bulletType="1", start=b.start, bulletFormat="%s.", **kw))
                elif task:
                    out.append(ListFlowable(items, bulletType="bullet", start="", leftIndent=4, spaceAfter=4,
                                            bulletFontName=self.font))
                else:
                    out.append(ListFlowable(items, bulletType="bullet", start="•", **kw))
            elif b.kind == "table":
                ncols = max(len(r) for r in b.rows)
                data = []
                for ri, row in enumerate(b.rows):
                    cells = []
                    for ci in range(ncols):
                        spans = row[ci] if ci < len(row) else []
                        if ri < b.header_rows:
                            spans = [Span(**{**s.__dict__, "bold": True}) for s in spans]
                        a = b.aligns[ci] if ci < len(b.aligns) else None
                        cells.append(Paragraph(self._markup(spans), self._cell_style(a)))
                    data.append(cells)
                t = Table(data, colWidths=[width / ncols] * ncols, repeatRows=b.header_rows)
                cmds = [("GRID", (0, 0), (-1, -1), 0.5, c.HexColor("#CFCCC5")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
                if b.header_rows:
                    cmds.append(("BACKGROUND", (0, 0), (-1, b.header_rows - 1), c.HexColor("#F1F0EC")))
                t.setStyle(TableStyle(cmds))
                out += [t, Spacer(1, self.cfg.base_size * 0.8)]
            elif b.kind == "code":
                pre = Preformatted(b.text, self.st["code"])
                box = Table([[pre]], colWidths=[width])
                box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), c.HexColor("#F4F3F0")),
                                         ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                                         ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                out += [box, Spacer(1, self.cfg.base_size * 0.8)]
            elif b.kind == "quote":
                inner = self.flowables(b.children, width - 14, "quote")
                box = Table([[inner]], colWidths=[width])
                box.setStyle(TableStyle([("LINEBEFORE", (0, 0), (0, -1), 2.5, c.HexColor("#C8C4BC")),
                                         ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 0),
                                         ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
                out += [box, Spacer(1, self.cfg.base_size * 0.6)]
            elif b.kind == "hr":
                out.append(HRFlowable(width="100%", thickness=0.6, color=c.HexColor("#CFCCC5"),
                                      spaceBefore=4, spaceAfter=10))
            elif b.kind == "pagebreak":
                out.append(PageBreak())
        return out

    def build(self, blocks: List[Block], out) -> None:
        from reportlab.lib.pagesizes import A4, LETTER
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas as rl_canvas
        from reportlab.platypus import SimpleDocTemplate

        cfg = self.cfg
        page = LETTER if cfg.page_size.upper() == "LETTER" else A4
        m = cfg.margin_mm * mm
        doc = SimpleDocTemplate(out, pagesize=page, leftMargin=m, rightMargin=m, topMargin=m, bottomMargin=m,
                                title=cfg.title or "", author=cfg.author or "", creator="aidoc-export")
        font = self.font
        header = cfg.title if (cfg.header and cfg.title) else None
        numbers = cfg.page_numbers

        class NumberedCanvas(rl_canvas.Canvas):
            """Buffers pages and draws them at the end so the total page count ("3 / 12") can be printed."""

            def __init__(self, *a, **kw):
                super().__init__(*a, **kw)
                self._pages = []

            def showPage(self):
                self._pages.append(dict(self.__dict__))
                self._startPage()

            def save(self):
                total = len(self._pages)
                for state in self._pages:
                    self.__dict__.update(state)
                    self.setFont(font, 8)
                    self.setFillColorRGB(0.53, 0.53, 0.53)
                    w, h = page
                    if header:
                        self.drawRightString(w - m, h - m * 0.55, header)
                    if numbers:
                        self.drawCentredString(w / 2, m * 0.45, "%d / %d" % (self._pageNumber, total))
                    super().showPage()
                super().save()

        doc.build(self.flowables(blocks, doc.width), canvasmaker=NumberedCanvas)


# ===========================================================================
# Public API
# ===========================================================================

def export_docx(markdown: str, out: PathOrFile = None, config: Optional[ExportConfig] = None) -> Optional[bytes]:
    """Markdown -> DOCX. Returns bytes if out is None."""
    cfg = config or ExportConfig()
    w = _DocxWriter(cfg)
    w.write(_append_sources(parse_markdown(markdown), cfg))
    return w.save(out)


def export_pdf(markdown: str, out: PathOrFile = None, config: Optional[ExportConfig] = None) -> Optional[bytes]:
    """Markdown -> PDF. Returns bytes if out is None. Requires a Korean TTF font (see find_korean_font)."""
    cfg = config or ExportConfig()
    w = _PdfWriter(cfg)
    blocks = _append_sources(parse_markdown(markdown), cfg)
    if out is None:
        buf = io.BytesIO()
        w.build(blocks, buf)
        return buf.getvalue()
    w.build(blocks, str(out) if isinstance(out, Path) else out)
    return None


# ===========================================================================
# CLI
# ===========================================================================

def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(prog="aidoc-export", description="Convert Markdown -> DOCX/PDF")
    ap.add_argument("input", help="Markdown file (- for stdin)")
    ap.add_argument("-o", "--output", required=True, help="Output file (.docx or .pdf)")
    ap.add_argument("--title")
    ap.add_argument("--author")
    ap.add_argument("--header", action="store_true", help="Show title in the header")
    ap.add_argument("--no-page-numbers", action="store_true")
    ap.add_argument("--page-size", default="A4", choices=["A4", "Letter"])
    ap.add_argument("--margin", type=float, default=20.0, help="Margin in mm (default 20)")
    ap.add_argument("--font", help="TTF for PDF body text")
    ap.add_argument("--font-bold", help="TTF for PDF bold text")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO if args.verbose else logging.WARNING, format="%(levelname)s %(message)s")

    text = sys.stdin.read() if args.input == "-" else Path(args.input).read_text(encoding="utf-8")
    cfg = ExportConfig(title=args.title, author=args.author, header=args.header,
                       page_numbers=not args.no_page_numbers, page_size=args.page_size, margin_mm=args.margin,
                       font_path=args.font, bold_font_path=args.font_bold)
    ext = Path(args.output).suffix.lower()
    try:
        if ext == ".docx":
            export_docx(text, args.output, cfg)
        elif ext == ".pdf":
            export_pdf(text, args.output, cfg)
        else:
            ap.error("Output extension must be .docx or .pdf")
    except FontNotFoundError as e:
        print("Error: %s" % e, file=sys.stderr)
        return 2
    print("Saved: %s" % args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
